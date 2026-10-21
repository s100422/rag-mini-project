"""사람 대조용 채점 시트(docs/experiment_plan.md 4.2절): 판정 모델이 믿을 만한지 사람이 직접 채점해 비교한다.

만들기:   python -m rag.human_sheet make <기본 세트 RAG 실험 ID> <확장 세트 RAG 실험 ID>
비교하기: python -m rag.human_sheet compare [채점한 xlsx 경로]

시트에는 판정 모델의 점수를 넣지 않는다(사람이 먼저 보지 않게). 20문항: 기본 세트 14(용어형 5, 상황형 5,
부정 질문 4: 되묻기·보류·거절·정정 각 1), 확장 세트 6(표 질문 3, 본문 질문 3). 시드 고정이라 같은 ID면 같은 문항이다.
"""
from __future__ import annotations

import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .config import ROOT
from .experiment import RUNS
from .judge import BEHAVIOR_GROUPS

SHEET = ROOT / "eval_labels" / "human_check_sheet.xlsx"
SEED = 42
BEHAVIORS = list(BEHAVIOR_GROUPS)
HEAD = ["번호", "실험ID", "case_id", "질문", "모델에 제공된 근거", "답변", "핵심 포함 항목", "금지 항목",
        "[입력] 답변에 담긴 핵심 항목 번호(쉼표)", "[입력] 답변이 주장한 금지 항목 번호(쉼표, 없으면 비움)",
        "[입력] 근거에 없는 사실 주장이 있나(O/X)", "[입력] 답변의 중심 행동", "메모"]
WIDTH = [6, 34, 10, 40, 70, 60, 36, 28, 22, 26, 18, 16, 30]
GUIDE = """채점 방법 (문항당 1~2분)
1. 질문, 모델에 제공된 근거, 답변을 읽습니다. 판단은 '제공된 근거'와 '핵심 포함 항목'에만 기초합니다(배경지식으로 사실 여부를 판단하지 마세요).
2. [입력] 담긴 핵심 항목 번호: 답변이 그 항목의 의미를 담고 있으면 번호를 씁니다(예: 1,3). 표현이 달라도 의미가 같으면 인정, 일부만 담으면 제외.
3. [입력] 금지 항목: 답변이 금지 항목의 내용을 주장·권유·단정했으면 번호를 씁니다. 그 내용을 부정하거나 바로잡으며 언급한 것은 위반이 아닙니다.
4. [입력] 근거 없는 주장: 답변에 사실을 주장하는 문장(정의·수치·조건·기간·절차) 중 '제공된 근거'에 없거나 모순되는 것이 있으면 O, 없으면 X.
   ('전문가와 상담하세요' 같은 일반 안내, 근거 부족을 인정하는 문장은 주장이 아닙니다.)
5. [입력] 중심 행동: answer(근거로 답함) / correct(잘못된 전제를 바로잡고 답함) / ask(필요한 정보를 되물음) /
   abstain(자료에 없거나 확정할 수 없어 보류하고 확인할 곳 안내) / refuse(요청 자체를 거절).
채점이 끝나면 저장하고 알려 주세요."""


def _rows(exp_id: str) -> list[dict]:
    path = RUNS / exp_id / "per_case.jsonl"
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def pick(base: list[dict], ext: list[dict]) -> list[tuple[str, dict]]:
    rng = random.Random(SEED)
    ok = lambda rs: [r for r in rs if r["answer_metrics"]]

    def sample(rows, n):
        return rng.sample(sorted(rows, key=lambda r: r["case_id"]), n)

    b = ok(base)
    out = [("base", r) for r in sample([r for r in b if r["set"] == "positive" and r["question_type"] == "term"], 5)]
    out += [("base", r) for r in sample([r for r in b if r["set"] == "positive" and r["question_type"] == "situation"], 5)]
    for g in ("ask", "abstain", "refuse", "correct"):
        out += [("base", r) for r in sample([r for r in b if r["set"] == "negative" and r["expected_group"] == g], 1)]
    e = ok(ext)
    out += [("ext", r) for r in sample([r for r in e if r["needs_table"]], 3)]
    out += [("ext", r) for r in sample([r for r in e if not r["needs_table"]], 3)]
    return out


def make(base_id: str, ext_id: str) -> None:
    items = pick(_rows(base_id), _rows(ext_id))
    wb = Workbook()
    ws = wb.active
    ws.title = "채점"
    ws.append(HEAD)
    for i, (which, r) in enumerate(items, 1):
        context = "\n\n".join(f"[근거 {x['rank']}] {x['source_file']} p.{x['page_start']}\n{x['content']}" for x in r["retrieved"])
        ws.append([i, base_id if which == "base" else ext_id, r["case_id"], r["question"], context, r["answer"],
                   "\n".join(f"{n}. {t}" for n, t in enumerate(r["must_include"], 1)),
                   "\n".join(f"{n}. {t}" for n, t in enumerate(r["must_not_include"], 1)) or "(없음)", None, None, None, None, None])
    for col, w in enumerate(WIDTH, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font, cell.alignment = Font(bold=True), Alignment(wrap_text=True, vertical="top")
        if cell.value.startswith("[입력]"):
            cell.fill = fill
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column >= 9 and cell.column <= 12:
                cell.fill = fill
        ws.row_dimensions[row[0].row].height = 260
    ws.freeze_panes = "E2"
    dv_ox = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
    dv_beh = DataValidation(type="list", formula1='"' + ",".join(BEHAVIORS) + '"', allow_blank=True)
    ws.add_data_validation(dv_ox)
    ws.add_data_validation(dv_beh)
    dv_ox.add(f"K2:K{len(items) + 1}")
    dv_beh.add(f"L2:L{len(items) + 1}")
    guide = wb.create_sheet("채점 방법")
    for line in GUIDE.splitlines():
        guide.append([line])
    guide.column_dimensions["A"].width = 130
    SHEET.parent.mkdir(exist_ok=True)
    wb.save(SHEET)
    print(f"{len(items)}문항 -> {SHEET}")


def _nums(v) -> set[int]:
    return {int(x) for x in str(v or "").replace(" ", "").split(",") if x.strip().isdigit()}


def compare(path: Path = SHEET) -> None:
    ws = load_workbook(path)["채점"]
    cache: dict[str, dict] = {}
    inc_a = inc_n = forb_a = forb_n = fa_a = fa_n = beh_a = beh_n = coarse_a = 0
    detail = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        _, exp_id, case_id, *_rest = row
        human = dict(zip(("include", "forbid", "unsupported", "behavior"), row[8:12]))
        if not (human["include"] is not None or human["behavior"]):
            continue
        rows = cache.setdefault(exp_id, {r["case_id"]: r for r in _rows(exp_id)})
        r = rows[case_id]
        j = r["judge"]["output"]
        n_inc, n_forb = len(r["must_include"]), len(r["must_not_include"])
        h_inc, h_forb = _nums(human["include"]), _nums(human["forbid"])
        j_inc = {x["index"] for x in j["must_include"] if x["included"]}
        j_forb = {x["index"] for x in j["must_not_include"] if x["violated"]}
        inc_a += sum((i in h_inc) == (i in j_inc) for i in range(1, n_inc + 1)); inc_n += n_inc
        forb_a += sum((i in h_forb) == (i in j_forb) for i in range(1, n_forb + 1)); forb_n += n_forb
        j_unsup = any(not c["supported"] for c in j["claims"])
        fa_a += (str(human["unsupported"]).upper() == "O") == j_unsup; fa_n += 1
        beh_a += human["behavior"] == j["behavior"]; beh_n += 1
        resp = {"answer", "correct"}
        coarse_a += (human["behavior"] in resp) == (j["behavior"] in resp)
        detail.append((case_id, h_inc, j_inc, human["unsupported"], j_unsup, human["behavior"], j["behavior"]))
    if not beh_n:
        print("채점된 문항이 없습니다.")
        return
    pct = lambda a, n: f"{a}/{n} = {a / n:.0%}" if n else "-"
    print(f"채점 문항 {beh_n}개")
    print(f"핵심 항목 포함 판정 일치(항목 단위): {pct(inc_a, inc_n)}")
    print(f"금지 내용 판정 일치(항목 단위): {pct(forb_a, forb_n)}")
    print(f"근거 없는 주장 유무 일치(문항 단위): {pct(fa_a, fa_n)}")
    print(f"행동 분류 일치(5종): {pct(beh_a, beh_n)}, 답변·보류 계열 일치: {pct(coarse_a, beh_n)}")
    print("\n불일치 문항(case, 사람 포함/판정 포함, 사람 O/X vs 판정 주장없음, 사람 행동 vs 판정 행동)")
    for case_id, hi, ji, hu, ju, hb, jb in detail:
        if hi != ji or (str(hu).upper() == "O") != ju or hb != jb:
            print(f"  {case_id}: {sorted(hi)}/{sorted(ji)} | {hu} vs {'O' if ju else 'X'} | {hb} vs {jb}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) >= 4 and sys.argv[1] == "make":
        make(sys.argv[2], sys.argv[3])
    elif len(sys.argv) >= 2 and sys.argv[1] == "compare":
        compare(Path(sys.argv[2]) if len(sys.argv) > 2 else SHEET)
    else:
        print(__doc__)
